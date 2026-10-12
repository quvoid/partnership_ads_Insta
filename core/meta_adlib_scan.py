# -*- coding: utf-8 -*-
"""
================================================================================
META AD LIBRARY SCAN - every ad a Facebook page has run (public, no login)
================================================================================
Opens https://www.facebook.com/ads/library/?view_all_page_id=<id> in headless
Chromium, scrolls until no more cards load, and parses each ad card: library id,
active/inactive, start (and end) date, platforms, body text, CTA, link domain,
how many ads share the creative, and the snapshot URL. Cross-checks the page's
own "N results" count against what was captured so gaps are visible.

  python core/meta_adlib_scan.py --pages pages.json --xlsx OUT.xlsx
  pages.json: {"phoenix_millennium": {"page_id": "114293561582939", "name": "..."} , ...}
================================================================================
"""
import sys, os, re, json, time
from datetime import datetime, timezone
from typing import Dict, Any, List

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from playwright.sync_api import sync_playwright

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36")

EXTRACT_JS = r"""
() => {
  const PLAT = ["Facebook","Instagram","Audience Network","Messenger","Threads","WhatsApp"];
  const cards = [];
  // every element whose own text starts with "Library ID:" anchors one card
  const anchors = [...document.querySelectorAll('div,span')].filter(e => e.childElementCount === 0 && /^Library ID:\s*\d+/.test(e.innerText || ""));
  const seen = new Set();
  for (const a of anchors) {
    const id = (a.innerText.match(/Library ID:\s*(\d+)/) || [])[1];
    if (!id || seen.has(id)) continue;
    seen.add(id);
    // climb to the card: first ancestor whose text contains both the id and "Sponsored"
    let card = a;
    for (let i = 0; i < 14 && card.parentElement; i++) {
      card = card.parentElement;
      const tx = card.innerText || "";
      if (tx.includes("Sponsored") && tx.includes(id) && (tx.match(/Library ID:/g) || []).length === 1) break;
    }
    const tx = card.innerText || "";
    const status = /^\s*(Active|Inactive)\b/.test(tx) ? (tx.match(/^\s*(Active|Inactive)/) || [])[1] : ((tx.match(/\b(Active|Inactive)\b/) || [])[1] || "");
    const started = (tx.match(/Started running on\s*([^\n·]+)/) || [])[1]?.trim() || "";
    const ran = tx.match(/(\d{1,2} \w{3} \d{4})\s*-\s*(\d{1,2} \w{3} \d{4})/) || [];
    const using = (tx.match(/(\d+)\s+ads?\s+use this creative/) || [])[1] || "1";
    const adv = (tx.match(/See (?:ad|summary) details\n([^\n]+)\nSponsored/) || [])[1]?.trim() || "";
    let body = "";
    const sp = tx.split(/\nSponsored\n/);
    if (sp.length > 1) body = sp[1].split(/\n(?:0:00 \/|[A-Z0-9][A-Z0-9.-]+\.(?:COM|IN|NET|ORG|APP|LIVE|SHOP|IO)\n)/)[0].trim().slice(0, 700);
    const dom = (tx.match(/\n([A-Z0-9][A-Z0-9.-]+\.(?:COM|IN|CO\.IN|NET|ORG|APP|LIVE|SHOP|IO))\n/) || [])[1] || "";
    const cta = (tx.match(/\n(Learn more|Shop now|Book now|Sign up|Get offer|Contact us|Send message|Apply now|Get directions|Watch more|Download|Order now|Send WhatsApp message|Visit Instagram Profile|Get quote|Get tickets)\n/i) || [])[1] || "";
    const labels = [...card.querySelectorAll('[aria-label]')].map(e => e.getAttribute('aria-label'));
    const plats = PLAT.filter(p => labels.some(l => l && l.toLowerCase().startsWith(p.toLowerCase())));
    const hasVideo = !!card.querySelector('video') || /0:00 \//.test(tx);
    const img = card.querySelector('img[src*="scontent"]:not([src*="s60x60"]):not([src*="s100x100"])');
    cards.push({ library_id: id, status, started, ran_from: ran[1] || "", ran_to: ran[2] || "", ads_using_creative: using,
                 advertiser: adv, body, cta, link_domain: dom, platforms: plats.join(", "), media: hasVideo ? "Video" : (img ? "Image" : ""),
                 snapshot_url: "https://www.facebook.com/ads/library/?id=" + id });
  }
  const countText = (document.body.innerText.match(/([\d,]+\+?)\s+results?/) || [])[1] || "";
  return { count_text: countText, cards };
}
"""


def scan_page(page, page_id: str, log=print) -> Dict[str, Any]:
    url = (f"https://www.facebook.com/ads/library/?active_status=all&ad_type=all&country=IN"
           f"&view_all_page_id={page_id}&search_type=page&media_type=all")
    page.goto(url, wait_until="domcontentloaded", timeout=60000)
    page.wait_for_timeout(5000)
    last = -1
    for _ in range(60):
        page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
        page.wait_for_timeout(1500)
        n = page.evaluate("(document.body.innerText.match(/Library ID:/g) || []).length")
        if n == last:
            page.wait_for_timeout(2000)
            n2 = page.evaluate("(document.body.innerText.match(/Library ID:/g) || []).length")
            if n2 == n:
                break
        last = n
    res = page.evaluate(EXTRACT_JS)
    res["page_id"] = page_id
    res["scanned_at"] = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    res["url"] = url
    log(f"    page {page_id}: site says {res['count_text'] or '?'} results, captured {len(res['cards'])} cards")
    return res


def run(pages: Dict[str, Dict[str, str]], log=print) -> Dict[str, Any]:
    out = {}
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        ctx = browser.new_context(user_agent=UA, viewport={"width": 1400, "height": 1000}, locale="en-GB",
                                  extra_http_headers={"Accept-Language": "en-GB,en;q=0.9"})
        pg = ctx.new_page()
        for key, meta in pages.items():
            log(f"[{key}]")
            try:
                out[key] = scan_page(pg, meta["page_id"], log)
                out[key]["name"] = meta.get("name", key)
            except Exception as e:
                log(f"    ERROR {type(e).__name__}: {str(e)[:120]}")
                out[key] = {"page_id": meta["page_id"], "name": meta.get("name", key), "cards": [], "count_text": "", "error": str(e)[:200]}
            time.sleep(2)
        browser.close()
    return out


def export(results: Dict[str, Any], xlsx: str):
    import openpyxl
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter
    from core.kolkata_engine import clean_cell
    bold = Font(bold=True)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Summary"
    ws.cell(1, 1, "Meta Ad Library - ads run by each page (India), captured " + datetime.now().strftime("%d %b %Y")).font = bold
    hdr = ["Account", "Facebook Page ID", "Ad Library Link", "Results Shown By Meta", "Ads Captured", "Active Now", "Inactive",
           "Video Ads", "Image Ads", "Earliest Start", "Latest Start", "Platforms Seen"]
    for c, h in enumerate(hdr, 1):
        ws.cell(3, c, h).font = bold
    r = 4
    def _d(s):
        try: return datetime.strptime(s, "%d %b %Y")
        except Exception: return None
    for key, res in results.items():
        cards = res.get("cards", [])
        dates = [d for d in (_d(c["started"]) for c in cards) if d]
        plats = sorted({p for c in cards for p in c["platforms"].split(", ") if p})
        vals = [f"@{key}", res["page_id"], res.get("url", ""), res.get("count_text", ""), len(cards),
                sum(1 for c in cards if c["status"] == "Active"), sum(1 for c in cards if c["status"] == "Inactive"),
                sum(1 for c in cards if c["media"] == "Video"), sum(1 for c in cards if c["media"] == "Image"),
                min(dates).strftime("%Y-%m-%d") if dates else "", max(dates).strftime("%Y-%m-%d") if dates else "", ", ".join(plats)]
        for c, v in enumerate(vals, 1):
            ws.cell(r, c, clean_cell(v))
        r += 1
    for i, w in enumerate([24, 18, 60, 14, 12, 10, 10, 10, 10, 14, 14, 36], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    for key, res in results.items():
        t = wb.create_sheet(re.sub(r'[\[\]\*\?/\\:]', '', key)[:31])
        t.cell(1, 1, f"@{key} - Meta Ad Library, page id {res['page_id']}").font = bold
        cols = ["Library ID", "Status", "Started Running", "Ran From", "Ran To", "Media", "Platforms", "Ads Using This Creative",
                "CTA", "Link Domain", "Ad Text", "Snapshot URL"]
        for c, h in enumerate(cols, 1):
            t.cell(3, c, h).font = bold
        cards = sorted(res.get("cards", []), key=lambda c: (c["status"] != "Active", -(_d(c["started"]) or datetime(1970, 1, 1)).timestamp()))
        for i, c in enumerate(cards, 4):
            for j, v in enumerate([c["library_id"], c["status"], c["started"], c["ran_from"], c["ran_to"], c["media"], c["platforms"],
                                   c["ads_using_creative"], c["cta"], c["link_domain"], c["body"], c["snapshot_url"]], 1):
                t.cell(i, j, clean_cell(v)).alignment = Alignment(vertical="top")
        for i, w in enumerate([18, 10, 16, 12, 12, 8, 26, 10, 18, 22, 90, 48], 1):
            t.column_dimensions[get_column_letter(i)].width = w

    m = wb.create_sheet("Method")
    notes = [("Source", "Meta Ad Library website, India, all ad types, per Facebook page (public data, no login)."),
             ("Page matching", "Each Instagram account was matched to its Facebook page through the Ad Library's own advertiser search, which shows the linked @instagram handle."),
             ("Status", "Active = running on the day captured. Inactive = has stopped; Meta keeps most non-political ads visible for a limited time after they stop, so old inactive ads may no longer appear."),
             ("Ads Using This Creative", "Meta groups identical creatives. The number is how many ad variants share that creative and text."),
             ("Ads Captured vs Results Shown", "Meta's count is per ad; the grid shows one card per creative group, so captured cards can be fewer than the results count without anything missing."),
             ("What is not here", "Spend, impressions and audience are only shown for political/issue ads. For commercial ads Meta publishes creative, dates, platforms and status only.")]
    m.cell(1, 1, "Field").font = bold; m.cell(1, 2, "Definition").font = bold
    for i, (k, v) in enumerate(notes, 2):
        m.cell(i, 1, k); m.cell(i, 2, v).alignment = Alignment(wrap_text=True, vertical="top")
    m.column_dimensions["A"].width = 30; m.column_dimensions["B"].width = 120
    path = xlsx if (os.sep in xlsx or "/" in xlsx) else os.path.join(BASE_DIR, "deliverables", xlsx)
    try:
        wb.save(path)
    except PermissionError:
        path = path.replace(".xlsx", f"_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"); wb.save(path)
    print(f"[EXCEL] {path}")
    return path


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    a = sys.argv[1:]
    pages = json.load(open(a[a.index("--pages") + 1], encoding="utf-8"))
    xlsx = a[a.index("--xlsx") + 1] if "--xlsx" in a else "Meta_AdLibrary.xlsx"
    res = run(pages)
    json.dump(res, open(os.path.join(BASE_DIR, "meta_adlib_cache.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    export(res, xlsx)
